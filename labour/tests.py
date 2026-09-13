from datetime import date
from decimal import Decimal
from io import BytesIO

from django.contrib.auth import get_user_model
from django.test import TestCase
from django.urls import reverse
from django.utils import timezone
from openpyxl import load_workbook

from .models import Labour, LabourAdvance, LabourDriverPayment, LabourExtraPayment, LabourRozi, LabourTripGroup
from .views import _labour_statement_for_period, _labour_type_label


class MistriStatementTest(TestCase):
    def setUp(self):
        user_model = get_user_model()
        self.user = user_model.objects.create_user(
            username='mistri-statement-tester',
            password='test-password-123',
        )
        self.client.force_login(self.user)
        self.mistri = Labour.objects.create(
            name='Raju Bhau',
            category='MISTRI',
            sub_category='MISTRI',
            base_daily_rate=Decimal('500.00'),
            is_active=True,
            status='ACTIVE',
        )
        self.day_full = date(2026, 9, 10)
        self.day_half = date(2026, 9, 9)
        LabourRozi.objects.create(
            labour=self.mistri, date=self.day_full, day_type='FULL',
        )
        LabourRozi.objects.create(
            labour=self.mistri, date=self.day_half, day_type='HALF',
        )
        LabourExtraPayment.objects.create(
            labour=self.mistri, date=self.day_full, amount=Decimal('100.00'),
        )
        LabourAdvance.objects.create(
            labour=self.mistri, date=self.day_full, amount=Decimal('200.00'),
        )
        self.period_start = date(2026, 9, 1)
        self.period_end = date(2026, 9, 14)

    def test_statement_has_rozi_breakdown_per_day(self):
        st = _labour_statement_for_period(
            self.mistri, self.period_start, self.period_end,
        )

        rozi = st['rozi_by_date'][self.day_full]
        self.assertEqual(rozi['day_type'], 'FULL')
        self.assertEqual(rozi['day_label'], 'Full Day')
        self.assertEqual(rozi['rate'], Decimal('800'))
        self.assertEqual(rozi['amount'], Decimal('800'))

        half = st['rozi_by_date'][self.day_half]
        self.assertEqual(half['day_label'], 'Half Day')
        self.assertEqual(half['amount'], Decimal('400'))

        self.assertEqual(st['rozi_total'], Decimal('1200'))

    def test_type_label_shows_mistri_roles(self):
        self.assertEqual(_labour_type_label(self.mistri), 'Mistri')
        helper = Labour.objects.create(
            name='Helper Lal',
            category='MISTRI',
            sub_category='HELPER',
            base_daily_rate=Decimal('500.00'),
            is_active=True,
            status='ACTIVE',
        )
        self.assertEqual(_labour_type_label(helper), 'Mistri Helper')
        driver = Labour.objects.create(
            name='Gaju Bhau',
            category='HYVA_DRIVER',
            is_active=True,
            status='ACTIVE',
            is_driver=True,
        )
        self.assertEqual(_labour_type_label(driver), 'Hyva Driver')

    def test_pdf_export_works_for_mistri(self):
        response = self.client.get(
            reverse('labour:statement_export', args=[self.mistri.id]),
            {'from_date': '2026-09-01', 'to_date': '2026-09-14'},
        )

        self.assertEqual(response.status_code, 200)
        self.assertEqual(response['Content-Type'], 'application/pdf')

    def test_excel_export_shows_day_type_rows_for_mistri(self):
        response = self.client.get(
            reverse('labour:statement_export', args=[self.mistri.id]),
            {
                'from_date': '2026-09-01',
                'to_date': '2026-09-14',
                'export': 'excel',
            },
        )
        self.assertEqual(response.status_code, 200)
        workbook = load_workbook(filename=BytesIO(response.content), read_only=True)
        rows = [
            [cell.value for cell in row]
            for row in workbook.active.iter_rows()
        ]
        values = [cell for row in rows for cell in row]
        self.assertIn('Day Type', values)
        self.assertIn('Full Day', values)
        self.assertIn('Half Day', values)
        # Rate ek baar top par, har line me nahi.
        self.assertNotIn('Rate', values)
        self.assertTrue(
            any(
                isinstance(cell, str)
                and 'Full Day' in cell
                and 'Half Day' in cell
                and 'Overtime' in cell
                for cell in values
            ),
            'rate info line missing from mistri sheet',
        )
        # TOTAL net hona chahiye: Rozi + Extra - Advance = 1200 + 100 - 200.
        header = next(
            row for row in rows
            if 'Day Type' in row
        )
        total_row = next(row for row in rows if 'TOTAL' in row)
        total_col = header.index('Total ₹')
        rozi_col = header.index('Rozi ₹')
        self.assertEqual(total_row[rozi_col], 1200)
        self.assertEqual(total_row[total_col], 1100)

    def test_excel_export_dates_ascending_for_mistri(self):
        response = self.client.get(
            reverse('labour:statement_export', args=[self.mistri.id]),
            {
                'from_date': '2026-09-01',
                'to_date': '2026-09-14',
                'export': 'excel',
            },
        )

        self.assertEqual(response.status_code, 200)
        workbook = load_workbook(filename=BytesIO(response.content), read_only=True)
        rows = [
            [cell.value for cell in row]
            for row in workbook.active.iter_rows()
        ]
        header_idx = next(
            idx for idx, row in enumerate(rows)
            if 'Day Type' in row
        )
        date_col = rows[header_idx].index('Date')
        data_dates = [
            row[date_col] for row in rows[header_idx + 1:]
            if row[date_col] and row[date_col] != 'TOTAL'
        ]
        self.assertEqual(data_dates, ['09-Sep-2026', '10-Sep-2026'])

    def test_pdf_mistri_rows_ascending(self):
        from core.pdf_utils import get_pdf_styles, get_registered_font
        from .views import _rozi_entries_data

        st = _labour_statement_for_period(
            self.mistri, self.period_start, self.period_end,
        )
        styles = get_pdf_styles(get_registered_font())
        data = _rozi_entries_data(st, styles)
        dates = [
            row[0].text for row in data[1:-1]
        ]
        self.assertEqual(dates, ['09-Sep-2026', '10-Sep-2026'])


class HyvaStatementTest(TestCase):
    def setUp(self):
        user_model = get_user_model()
        self.user = user_model.objects.create_user(
            username='hyva-statement-tester',
            password='test-password-123',
        )
        self.client.force_login(self.user)
        self.driver = Labour.objects.create(
            name='Gaju Bhau',
            category='HYVA_DRIVER',
            is_active=True,
            status='ACTIVE',
            is_driver=True,
        )
        self.group_later = LabourTripGroup.objects.create(
            date=date(2026, 9, 10),
            trip_count=2,
            load_type='FLYASH_HYVA',
        )
        self.group_later.labourers.add(self.driver)
        self.group_earlier = LabourTripGroup.objects.create(
            date=date(2026, 9, 8),
            trip_count=3,
            load_type='WHITE_HYVA',
        )
        self.group_earlier.labourers.add(self.driver)
        LabourExtraPayment.objects.create(
            labour=self.driver, date=date(2026, 9, 9), amount=Decimal('500.00'),
        )
        LabourDriverPayment.objects.create(
            labour=self.driver,
            period_start=date(2026, 9, 1),
            period_end=date(2026, 9, 14),
            amount=Decimal('12000.00'),
        )
        self.period_start = date(2026, 9, 1)
        self.period_end = date(2026, 9, 14)

    def _statement(self):
        return _labour_statement_for_period(
            self.driver, self.period_start, self.period_end,
        )

    def test_hyva_rate_info_shows_each_rate_once(self):
        from .views import _hyva_rate_info

        info = _hyva_rate_info(self._statement())

        self.assertIn('Fly Ash Hyva: ₹100', info)
        self.assertIn('White Sand Hyva: ₹200', info)
        self.assertEqual(info.count('₹100'), 1)
        self.assertEqual(info.count('₹200'), 1)

    def test_rate_pairs_split_two_columns(self):
        from .views import _hyva_rate_pairs, _rate_box_rows

        pairs = _hyva_rate_pairs(self._statement())
        rows = _rate_box_rows(pairs)

        self.assertEqual(len(pairs), 2)
        self.assertEqual(rows, [[pairs[0], pairs[1]]])

    def test_mistri_rate_pairs(self):
        from .views import _mistri_rate_pairs, _rate_box_rows

        mistri = Labour.objects.create(
            name='Rate Mistri',
            category='MISTRI',
            sub_category='MISTRI',
            base_daily_rate=Decimal('500.00'),
            is_active=True,
            status='ACTIVE',
        )
        pairs = _mistri_rate_pairs(mistri)
        rows = _rate_box_rows(pairs)

        self.assertEqual(
            [label for label, _rate in pairs],
            ['Full Day', 'Half Day', 'Overtime'],
        )
        self.assertEqual(len(rows), 2)
        self.assertEqual(len(rows[0]), 2)
        self.assertEqual(len(rows[1]), 1)

    def test_compact_cards_wrap_two_per_row(self):
        from core.pdf_utils import build_summary_cards

        table = build_summary_cards(
            [
                {'label': f'C{i}', 'value': f'₹{i}.00'}
                for i in range(8)
            ],
            columns=2,
            value_size=10.5,
            pad=5,
        )

        self.assertEqual(len(table._cellvalues), 4)
        self.assertTrue(
            all(len(row) == 2 for row in table._cellvalues)
        )

    def test_ordered_day_entries_interleave_ascending(self):
        from .views import _ordered_day_entries

        entries = _ordered_day_entries(self._statement())

        self.assertEqual(
            [
                (kind, item.date if kind == 'group' else item['date'])
                for kind, item in entries
            ],
            [
                ('group', date(2026, 9, 8)),
                ('extra', date(2026, 9, 9)),
                ('group', date(2026, 9, 10)),
            ],
        )

    def test_category_summary_groups_by_load(self):
        from .views import _category_summary

        summary, grand_trips, grand_amount = _category_summary(self._statement())

        self.assertEqual(
            summary,
            [
                ('Hyva (Fly Ash Hyva)', 2, Decimal('200')),
                ('Hyva (White Sand Hyva)', 3, Decimal('600')),
            ],
        )
        self.assertEqual(grand_trips, 5)
        self.assertEqual(grand_amount, Decimal('800'))

    def test_excel_hyva_layout(self):
        response = self.client.get(
            reverse('labour:statement_export', args=[self.driver.id]),
            {
                'from_date': '2026-09-01',
                'to_date': '2026-09-14',
                'export': 'excel',
            },
        )

        self.assertEqual(response.status_code, 200)
        workbook = load_workbook(filename=BytesIO(response.content), read_only=True)
        rows = [
            [cell.value for cell in row]
            for row in workbook.active.iter_rows()
        ]
        values = [cell for row in rows for cell in row]
        # Rate ek baar top par.
        self.assertTrue(
            any(
                isinstance(cell, str) and 'Fly Ash Hyva' in cell and '₹200' in cell
                for cell in values
            ),
            'hyva rate info line missing',
        )
        # Bhatta row date-position par (9th, groups ke beech me).
        entries_header_idx = next(
            idx for idx, row in enumerate(rows)
            if 'Description' in row
        )
        summary_idx = next(
            idx for idx, row in enumerate(rows)
            if 'WORK SUMMARY' in row
        )
        date_col = rows[entries_header_idx].index('Date')
        entry_dates = [
            row[date_col]
            for row in rows[entries_header_idx + 1:summary_idx]
            if row[date_col]
        ]
        self.assertEqual(
            entry_dates,
            ['08-Sep-2026', '09-Sep-2026', '10-Sep-2026'],
        )
        # Entries table ka TOTAL row hata diya.
        self.assertNotIn('TOTAL', entry_dates)
        # Category summary + grand total.
        summary_values = [
            cell for row in rows[summary_idx:] for cell in row
        ]
        self.assertIn(800, summary_values)

    def test_payment_summary_rows_split_month_payment(self):
        from .views import _payment_summary_rows

        rows = _payment_summary_rows(self._statement(), 'Total Bhatta')

        self.assertEqual(
            [label for label, _value in rows],
            [
                'Total Earning',
                'Total Bhatta',
                'Month Payment',
                None,
                'Total Income Earned',
                None,
                'Total Advanced',
                'Old Balance',
                'Final Payment',
            ],
        )
        values = dict(
            (label, value) for label, value in rows if label is not None
        )
        self.assertEqual(values['Total Earning'], Decimal('800'))
        self.assertEqual(values['Total Bhatta'], Decimal('500'))
        self.assertEqual(values['Month Payment'], Decimal('12000'))
        self.assertEqual(values['Total Income Earned'], Decimal('13300'))

    def test_payment_summary_without_month_payment(self):
        from .views import _payment_summary_rows

        LabourDriverPayment.objects.all().delete()
        rows = _payment_summary_rows(self._statement(), 'Total Bhatta')

        self.assertEqual(
            [label for label, _value in rows],
            [
                'Total Earning',
                'Total Bhatta',
                None,
                'Total Income Earned',
                None,
                'Total Advanced',
                'Old Balance',
                'Final Payment',
            ],
        )
        values = dict(
            (label, value) for label, value in rows if label is not None
        )
        self.assertEqual(values['Total Income Earned'], Decimal('1300'))

    def test_excel_payment_summary_block(self):
        response = self.client.get(
            reverse('labour:statement_export', args=[self.driver.id]),
            {
                'from_date': '2026-09-01',
                'to_date': '2026-09-14',
                'export': 'excel',
            },
        )

        self.assertEqual(response.status_code, 200)
        workbook = load_workbook(filename=BytesIO(response.content), read_only=True)
        rows = [
            [cell.value for cell in row]
            for row in workbook.active.iter_rows()
        ]
        pay_idx = next(
            idx for idx, row in enumerate(rows)
            if 'PAYMENT SUMMARY' in row
        )
        work_idx = next(
            idx for idx, row in enumerate(rows)
            if 'WORK SUMMARY' in row
        )
        # Payment Summary neeche (stacked), alag section me.
        self.assertGreater(pay_idx, work_idx)
        pay_col = rows[pay_idx].index('PAYMENT SUMMARY')
        self.assertEqual(pay_col, 0)
        pay_labels = [
            row[pay_col] for row in rows[pay_idx + 1:]
            if row[pay_col]
        ]
        self.assertEqual(
            pay_labels,
            [
                'Description',
                'Total Earning',
                'Total Bhatta',
                'Month Payment',
                'Total Income Earned',
                'Total Advanced',
                'Old Balance',
                'Final Payment',
            ],
        )
        pay_values = {
            row[pay_col]: row[pay_col + 1] for row in rows[pay_idx + 1:]
            if row[pay_col] and row[pay_col] != 'Description'
        }
        self.assertEqual(pay_values['Month Payment'], 12000)
        self.assertEqual(pay_values['Total Income Earned'], 13300)

        # Total Advanced amount red me.
        cells = list(workbook.active.iter_rows())
        advanced_row = next(
            row for row in cells
            if row[pay_col].value == 'Total Advanced'
        )
        self.assertEqual(
            advanced_row[pay_col + 1].font.color.rgb, '00DC2626'
        )

        # Final Payment bold.
        cells = list(workbook.active.iter_rows())
        final_row = next(
            row for row in cells
            if row[pay_col].value == 'Final Payment'
        )
        self.assertTrue(final_row[pay_col].font.bold)
        self.assertTrue(final_row[pay_col + 1].font.bold)

    def test_summary_sections_stay_together(self):
        from reportlab.platypus import KeepTogether

        from core.pdf_utils import get_pdf_styles, get_registered_font
        from .views import _summary_flowables

        st = self._statement()
        styles = get_pdf_styles(get_registered_font())
        blocks = _summary_flowables(st, self.driver, styles, get_registered_font())

        # Work + Payment blocks, dono page-split se protected.
        self.assertEqual(len(blocks), 2)
        self.assertTrue(
            all(isinstance(block, KeepTogether) for block in blocks)
        )


class TractorStatementTest(TestCase):
    def setUp(self):
        user_model = get_user_model()
        self.user = user_model.objects.create_user(
            username='tractor-statement-tester',
            password='test-password-123',
        )
        self.client.force_login(self.user)
        self.labour = Labour.objects.create(
            name='Ankush Bhau',
            category='TRACTOR',
            base_daily_rate=Decimal('450.00'),
            is_active=True,
            status='ACTIVE',
            is_driver=True,
        )
        group = LabourTripGroup.objects.create(
            date=date(2026, 9, 10),
            trip_count=6,
            rate_per_trip=Decimal('450.00'),
        )
        group.labourers.add(self.labour)
        LabourExtraPayment.objects.create(
            labour=self.labour, date=date(2026, 9, 10), amount=Decimal('300.00'),
        )
        LabourAdvance.objects.create(
            labour=self.labour, date=date(2026, 9, 10), amount=Decimal('500.00'),
        )
        self.period_start = date(2026, 9, 1)
        self.period_end = date(2026, 9, 14)

    def _statement(self):
        return _labour_statement_for_period(
            self.labour, self.period_start, self.period_end,
        )

    def test_tractor_entries_have_no_trip_columns(self):
        from core.pdf_utils import get_pdf_styles, get_registered_font
        from .views import _rozi_entries_data

        st = self._statement()
        styles = get_pdf_styles(get_registered_font())
        data = _rozi_entries_data(st, styles, show_day_type=False)

        headers = [cell.text for cell in data[0]]
        self.assertEqual(
            headers,
            ['<b>Date</b>', '<b>Rozi (₹)</b>', '<b>Extra (₹)</b>', '<b>Advance ₹</b>', '<b>Total ₹</b>'],
        )
        # Day earning 2700 + extra 300 - advance 500 = net 2500.
        totals = [cell.text for cell in data[-1]]
        self.assertTrue(any('₹2,700.00' in cell for cell in totals))
        self.assertTrue(any('₹2,500.00' in cell for cell in totals))

    def test_tractor_rate_info_once_on_top(self):
        from .views import _tractor_rate_info

        info = _tractor_rate_info(self._statement())

        self.assertIn('₹450', info)
        self.assertEqual(info.count('₹450'), 1)

    def test_excel_tractor_layout(self):
        response = self.client.get(
            reverse('labour:statement_export', args=[self.labour.id]),
            {
                'from_date': '2026-09-01',
                'to_date': '2026-09-14',
                'export': 'excel',
            },
        )

        self.assertEqual(response.status_code, 200)
        workbook = load_workbook(filename=BytesIO(response.content), read_only=True)
        rows = [
            [cell.value for cell in row]
            for row in workbook.active.iter_rows()
        ]
        values = [cell for row in rows for cell in row]
        # Trip columns nahi, Work Summary nahi.
        self.assertNotIn('Trips', values)
        self.assertNotIn('WORK SUMMARY', values)
        header = next(row for row in rows if 'Rozi ₹' in row)
        self.assertEqual(
            [cell for cell in header if cell],
            ['Date', 'Rozi ₹', 'Extra ₹', 'Advance ₹', 'Total ₹'],
        )
        total_row = next(row for row in rows if 'TOTAL' in row)
        self.assertEqual(total_row[header.index('Total ₹')], 2500)
